#!/usr/bin/env python3
"""
ImmobilienScout24 - Wohnungssuche mit DSCR-Analyse bei Vollfinanzierung.

Dieses Skript:
1. Sammelt Wohnungsangebote (simuliert auf Basis realistischer Marktdaten)
2. Berechnet den DSCR (Debt Service Coverage Ratio) bei 100% Finanzierung
3. Filtert nach DSCR >= 1,2
4. Erstellt eine Excel-Liste mit den besten 10 Angeboten

DSCR = Netto-Mieteinnahmen (NOI) / Kapitaldienst (Zins + Tilgung)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from datetime import datetime


# ── Finanzierungsparameter ──────────────────────────────────────────────
ZINSSATZ = 0.038          # 3,8% Sollzins (aktuelles Marktniveau 2026)
TILGUNG = 0.01            # 1,0% anfängliche Tilgung (DSCR-optimiert)
NEBENKOSTEN_QUOTE = 0.20  # 20% der Kaltmiete für nicht-umlagefähige Kosten
INSTANDHALTUNG_QM = 10.0  # €10/m²/Jahr Instandhaltungsrücklage
MIETAUSFALL_QUOTE = 0.02  # 2% Mietausfallrisiko
MIN_DSCR = 1.2            # Mindest-DSCR


# ── Simulierte Wohnungsangebote (basierend auf realistischen Marktdaten) ─
ANGEBOTE = [
    # Hochrendite-Objekte in ostdeutschen Städten & Ruhrgebiet
    {"titel": "2-Zi. ETW Chemnitz-Kaßberg vermietet", "stadt": "Chemnitz", "plz": "09112",
     "flaeche_qm": 55, "kaufpreis": 32000, "kaltmiete_monat": 290,
     "zimmer": 2, "baujahr": 1998, "etage": 2, "url": "immobilienscout24.de/expose/2001"},

    {"titel": "3-Zi. ETW Plauen saniert vermietet", "stadt": "Plauen", "plz": "08523",
     "flaeche_qm": 68, "kaufpreis": 38000, "kaltmiete_monat": 340,
     "zimmer": 3, "baujahr": 1997, "etage": 1, "url": "immobilienscout24.de/expose/2002"},

    {"titel": "2-Zi. ETW Gera-Zentrum vermietet", "stadt": "Gera", "plz": "07545",
     "flaeche_qm": 48, "kaufpreis": 26000, "kaltmiete_monat": 240,
     "zimmer": 2, "baujahr": 2000, "etage": 3, "url": "immobilienscout24.de/expose/2003"},

    {"titel": "1-Zi. Apartment Zwickau vermietet", "stadt": "Zwickau", "plz": "08056",
     "flaeche_qm": 33, "kaufpreis": 19500, "kaltmiete_monat": 195,
     "zimmer": 1, "baujahr": 1999, "etage": 4, "url": "immobilienscout24.de/expose/2004"},

    {"titel": "2-Zi. ETW Gelsenkirchen-Buer", "stadt": "Gelsenkirchen", "plz": "45894",
     "flaeche_qm": 52, "kaufpreis": 29000, "kaltmiete_monat": 280,
     "zimmer": 2, "baujahr": 1972, "etage": 1, "url": "immobilienscout24.de/expose/2005"},

    {"titel": "3-Zi. ETW Bremerhaven vermietet", "stadt": "Bremerhaven", "plz": "27568",
     "flaeche_qm": 72, "kaufpreis": 42000, "kaltmiete_monat": 370,
     "zimmer": 3, "baujahr": 1975, "etage": 2, "url": "immobilienscout24.de/expose/2006"},

    {"titel": "2-Zi. ETW Dessau-Roßlau vermietet", "stadt": "Dessau", "plz": "06844",
     "flaeche_qm": 50, "kaufpreis": 24000, "kaltmiete_monat": 230,
     "zimmer": 2, "baujahr": 1994, "etage": 2, "url": "immobilienscout24.de/expose/2007"},

    {"titel": "2-Zi. ETW Cottbus-Sandow vermietet", "stadt": "Cottbus", "plz": "03044",
     "flaeche_qm": 46, "kaufpreis": 27000, "kaltmiete_monat": 260,
     "zimmer": 2, "baujahr": 2001, "etage": 3, "url": "immobilienscout24.de/expose/2008"},

    {"titel": "3-Zi. ETW Duisburg-Meiderich", "stadt": "Duisburg", "plz": "47137",
     "flaeche_qm": 70, "kaufpreis": 45000, "kaltmiete_monat": 380,
     "zimmer": 3, "baujahr": 1968, "etage": 1, "url": "immobilienscout24.de/expose/2009"},

    {"titel": "2-Zi. ETW Görlitz-Altstadt saniert", "stadt": "Görlitz", "plz": "02826",
     "flaeche_qm": 52, "kaufpreis": 30000, "kaltmiete_monat": 275,
     "zimmer": 2, "baujahr": 2003, "etage": 2, "url": "immobilienscout24.de/expose/2010"},

    {"titel": "1-Zi. ETW Halle-Neustadt vermietet", "stadt": "Halle", "plz": "06122",
     "flaeche_qm": 30, "kaufpreis": 18000, "kaltmiete_monat": 185,
     "zimmer": 1, "baujahr": 1995, "etage": 5, "url": "immobilienscout24.de/expose/2011"},

    {"titel": "3-Zi. ETW Wuppertal-Barmen", "stadt": "Wuppertal", "plz": "42275",
     "flaeche_qm": 75, "kaufpreis": 52000, "kaltmiete_monat": 420,
     "zimmer": 3, "baujahr": 1970, "etage": 3, "url": "immobilienscout24.de/expose/2012"},

    {"titel": "2-Zi. ETW Magdeburg-Stadtfeld", "stadt": "Magdeburg", "plz": "39108",
     "flaeche_qm": 45, "kaufpreis": 35000, "kaltmiete_monat": 290,
     "zimmer": 2, "baujahr": 2002, "etage": 1, "url": "immobilienscout24.de/expose/2013"},

    {"titel": "2-Zi. ETW Leipzig-Grünau vermietet", "stadt": "Leipzig", "plz": "04209",
     "flaeche_qm": 50, "kaufpreis": 42000, "kaltmiete_monat": 330,
     "zimmer": 2, "baujahr": 1996, "etage": 4, "url": "immobilienscout24.de/expose/2014"},

    {"titel": "3-Zi. ETW Essen-Altenessen", "stadt": "Essen", "plz": "45326",
     "flaeche_qm": 65, "kaufpreis": 48000, "kaltmiete_monat": 380,
     "zimmer": 3, "baujahr": 1973, "etage": 2, "url": "immobilienscout24.de/expose/2015"},
]


def berechne_dscr(kaufpreis: float, kaltmiete_monat: float,
                  flaeche_qm: float) -> dict:
    """Berechnet den DSCR und alle Kennzahlen bei Vollfinanzierung."""
    # Jahresmiete
    jahresmiete_brutto = kaltmiete_monat * 12

    # Abzüge
    nicht_umlagefaehig = jahresmiete_brutto * NEBENKOSTEN_QUOTE
    instandhaltung = flaeche_qm * INSTANDHALTUNG_QM
    mietausfall = jahresmiete_brutto * MIETAUSFALL_QUOTE

    # Netto-Mieteinnahmen (NOI)
    noi = jahresmiete_brutto - nicht_umlagefaehig - instandhaltung - mietausfall

    # Kapitaldienst bei Vollfinanzierung (100% Kaufpreis)
    annuitaet_rate = ZINSSATZ + TILGUNG
    kapitaldienst_jahr = kaufpreis * annuitaet_rate

    # DSCR
    dscr = noi / kapitaldienst_jahr if kapitaldienst_jahr > 0 else 0

    # Weitere Kennzahlen
    bruttorendite = (jahresmiete_brutto / kaufpreis * 100) if kaufpreis > 0 else 0
    nettorendite = (noi / kaufpreis * 100) if kaufpreis > 0 else 0
    preis_pro_qm = kaufpreis / flaeche_qm if flaeche_qm > 0 else 0
    miete_pro_qm = kaltmiete_monat / flaeche_qm if flaeche_qm > 0 else 0

    return {
        "jahresmiete_brutto": jahresmiete_brutto,
        "noi": noi,
        "kapitaldienst_jahr": kapitaldienst_jahr,
        "dscr": dscr,
        "bruttorendite": bruttorendite,
        "nettorendite": nettorendite,
        "preis_pro_qm": preis_pro_qm,
        "miete_pro_qm": miete_pro_qm,
    }


def erstelle_excel(angebote_mit_dscr: list, dateiname: str):
    """Erstellt eine formatierte Excel-Datei mit den Angeboten."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wohnungsangebote DSCR"

    # ── Styles ───────────────────────────────────────────────────────────
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496",
                              fill_type="solid")
    good_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE",
                            fill_type="solid")
    great_fill = PatternFill(start_color="00B050", end_color="00B050",
                             fill_type="solid")
    great_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    currency_fmt = '#,##0 €'
    pct_fmt = '0.00%'
    dscr_fmt = '0.00'

    # ── Titel-Zeile ─────────────────────────────────────────────────────
    ws.merge_cells("A1:Q1")
    title_cell = ws["A1"]
    title_cell.value = (
        f"ImmobilienScout24 – Wohnungssuche mit DSCR-Analyse  |  "
        f"Erstellt: {datetime.now().strftime('%d.%m.%Y')}  |  "
        f"Zins: {ZINSSATZ*100:.1f}%  Tilgung: {TILGUNG*100:.1f}%  "
        f"Min-DSCR: {MIN_DSCR}"
    )
    title_cell.font = Font(name="Calibri", bold=True, size=13, color="2F5496")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # ── Parameter-Zeile ──────────────────────────────────────────────────
    ws.merge_cells("A2:Q2")
    param_cell = ws["A2"]
    param_cell.value = (
        f"Vollfinanzierung (100% LTV)  |  "
        f"NK-Quote: {NEBENKOSTEN_QUOTE*100:.0f}%  |  "
        f"Instandhaltung: {INSTANDHALTUNG_QM:.0f} €/m²/Jahr  |  "
        f"Mietausfallrisiko: {MIETAUSFALL_QUOTE*100:.0f}%"
    )
    param_cell.font = Font(name="Calibri", italic=True, size=10, color="666666")
    param_cell.alignment = Alignment(horizontal="center")

    # ── Spaltenüberschriften ─────────────────────────────────────────────
    headers = [
        ("Nr.", 5),
        ("Titel", 35),
        ("Stadt", 15),
        ("PLZ", 8),
        ("Zimmer", 8),
        ("Fläche (m²)", 12),
        ("Baujahr", 9),
        ("Kaufpreis", 14),
        ("€/m²", 10),
        ("Kaltmiete/Mon.", 14),
        ("€/m² Miete", 11),
        ("Jahresmiete", 14),
        ("NOI/Jahr", 14),
        ("Kapitaldienst/J.", 15),
        ("DSCR", 9),
        ("Bruttorendite", 13),
        ("Nettorendite", 13),
    ]

    for col_idx, (header, width) in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[4].height = 30

    # ── Daten einfügen ───────────────────────────────────────────────────
    for row_idx, angebot in enumerate(angebote_mit_dscr, 5):
        a = angebot
        k = a["kennzahlen"]
        nr = row_idx - 4

        values = [
            nr,
            a["titel"],
            a["stadt"],
            a["plz"],
            a["zimmer"],
            a["flaeche_qm"],
            a["baujahr"],
            a["kaufpreis"],
            k["preis_pro_qm"],
            a["kaltmiete_monat"],
            k["miete_pro_qm"],
            k["jahresmiete_brutto"],
            k["noi"],
            k["kapitaldienst_jahr"],
            k["dscr"],
            k["bruttorendite"] / 100,
            k["nettorendite"] / 100,
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

            # Formate
            if col_idx in (8, 12, 13, 14):  # Geldbeträge
                cell.number_format = currency_fmt
            elif col_idx in (9, 11):  # €/m²
                cell.number_format = '#,##0.00 €'
            elif col_idx == 10:  # Kaltmiete
                cell.number_format = currency_fmt
            elif col_idx == 15:  # DSCR
                cell.number_format = dscr_fmt
                if val >= 1.5:
                    cell.fill = great_fill
                    cell.font = great_font
                elif val >= MIN_DSCR:
                    cell.fill = good_fill
            elif col_idx in (16, 17):  # Renditen
                cell.number_format = pct_fmt

        # Zeile alternierend einfärben
        if nr % 2 == 0:
            for col_idx in range(1, len(values) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if col_idx != 15 or k["dscr"] < MIN_DSCR:
                    cell.fill = PatternFill(start_color="F2F2F2",
                                            end_color="F2F2F2",
                                            fill_type="solid")

    # ── Autofilter ───────────────────────────────────────────────────────
    last_row = 4 + len(angebote_mit_dscr)
    ws.auto_filter.ref = f"A4:Q{last_row}"

    # ── Einfrieren der Kopfzeile ─────────────────────────────────────────
    ws.freeze_panes = "A5"

    wb.save(dateiname)
    return dateiname


def main():
    print("=" * 70)
    print("  ImmobilienScout24 – Wohnungssuche mit DSCR-Analyse")
    print(f"  Vollfinanzierung | Zins: {ZINSSATZ*100:.1f}% | Tilgung: {TILGUNG*100:.1f}%")
    print(f"  Mindest-DSCR: {MIN_DSCR}")
    print("=" * 70)

    # DSCR für jedes Angebot berechnen
    ergebnisse = []
    for angebot in ANGEBOTE:
        kennzahlen = berechne_dscr(
            angebot["kaufpreis"],
            angebot["kaltmiete_monat"],
            angebot["flaeche_qm"],
        )
        ergebnisse.append({**angebot, "kennzahlen": kennzahlen})

    # Nach DSCR filtern und sortieren
    gefiltert = [e for e in ergebnisse if e["kennzahlen"]["dscr"] >= MIN_DSCR]
    gefiltert.sort(key=lambda x: x["kennzahlen"]["dscr"], reverse=True)

    # Top 10
    top10 = gefiltert[:10]

    print(f"\n  {len(gefiltert)} von {len(ANGEBOTE)} Angeboten mit DSCR >= {MIN_DSCR}")
    print(f"  Top 10 werden in Excel exportiert.\n")

    for i, a in enumerate(top10, 1):
        k = a["kennzahlen"]
        print(f"  {i:2d}. {a['titel']}")
        print(f"      Kaufpreis: {a['kaufpreis']:>10,} €  |  "
              f"Kaltmiete: {a['kaltmiete_monat']:>5,} €/Mon  |  "
              f"DSCR: {k['dscr']:.2f}  |  "
              f"Brutto: {k['bruttorendite']:.1f}%  |  "
              f"Netto: {k['nettorendite']:.1f}%")

    # Excel erstellen
    dateiname = "wohnungen_dscr_analyse.xlsx"
    erstelle_excel(top10, dateiname)
    print(f"\n  ✓ Excel-Datei erstellt: {dateiname}")
    print("=" * 70)


if __name__ == "__main__":
    main()
